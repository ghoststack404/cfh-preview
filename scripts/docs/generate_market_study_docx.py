# -*- coding: utf-8 -*-
"""
Generador del Documento Independiente: Estudio de Mercado y Benchmarking Competitivo
para CoachingFromHome (CFH).
Produce: 'Estudio de Mercado.docx'
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_srs_helpers import (
    set_cell_background, set_cell_margins, set_table_borders,
    add_callout, add_code_block, format_table
)

def create_market_study_doc():
    doc = docx.Document()

    # Configuración de página (Márgenes de 1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("CoachingFromHome (CFH) • Estudio de Mercado y Benchmarking Competitivo • Confidencial")
        r_ft.font.name = "Calibri"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = RGBColor(6, 78, 59)  # Emerald 900
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(4, 120, 87)  # Emerald 700
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900
        return p

    def add_p(text, bold_prefix=None, space_after=5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Calibri"
            r_b.font.size = Pt(10)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(15, 23, 42)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Calibri"
            r_b.font.size = Pt(10)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(15, 23, 42)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # =========================================================================
    # PORTADA CORPORATIVA
    # =========================================================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(36)
    p_pre.paragraph_format.space_after = Pt(6)
    r_pre = p_pre.add_run("INTELIGENCIA COMPETITIVA, PRODUCTO & ARQUITECTURA UX/UI")
    r_pre.font.name = "Calibri"
    r_pre.font.size = Pt(11)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(4, 120, 87)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(10)
    r_title = p_title.add_run("ESTUDIO DE MERCADO EXHAUSTIVO Y BENCHMARKING COMPETITIVO")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(10, 37, 64)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(28)
    r_sub = p_sub.add_run(
        "Análisis Técnico y Funcional de Plataformas Líderes Globales (Preply, Superprof, Udemy, Domestika, MentorCruise, ADPList y Skool)\n"
        "Guía de Referencia para la Reestructuración, Adopción de Herramientas y Evolución de CoachingFromHome (CFH)"
    )
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    meta_table = doc.add_table(rows=1, cols=2)
    format_table(
        meta_table,
        col_widths=[2.5, 4.0],
        headers=["Atributo del Estudio", "Detalle / Especificación"],
        rows_data=[
            ["Proyecto Objetivo", "CoachingFromHome (CFH) - https://coachingfromhome.com.mx"],
            ["Título del Entregable", "Estudio de Mercado y Benchmarking Funcional para Reestructuración"],
            ["Plataformas Analizadas", "Preply, Superprof, Udemy, Domestika, MentorCruise, ADPList, Skool"],
            ["Áreas Evaluadas", "UI/UX, Conversión, Aula Virtual, IA Post-Sesión, Anti-Desintermediación, Comunidad"],
            ["Fecha de Elaboración", "Octubre 2026"],
            ["Autoría", "Arquitectura de Software Principal & Analista de Producto Senior"],
            ["Propósito", "Servir como catálogo de referencia práctica para copiar, adaptar y superar las mejores herramientas del mercado."]
        ]
    )

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 1: RESUMEN EJECUTIVO Y POSICIONAMIENTO ESTRATÉGICO
    # =========================================================================
    add_h1("1. Resumen Ejecutivo y Posicionamiento de CoachingFromHome")
    add_p(
        "El mercado global de EdTech, mentoría personalizada 1 a 1 y aprendizaje basado en comunidades ha evolucionado "
        "hacia ecosistemas integrados donde la experiencia de usuario (UX), la inteligencia artificial aplicada y la retención "
        "transaccional determinan el éxito comercial de una plataforma. CoachingFromHome (CFH) posee una ventaja competitiva "
        "única al reunir bajo un mismo techo cuatro pilares que en el mercado suelen encontrarse fragmentados en distintas aplicaciones:"
    )
    add_bullet("1. Marketplace de Sesiones 1 a 1 en Vivo (Competidores directos: Preply, Superprof).", "Pilar 1: ")
    add_bullet("2. Catálogo y Reproductor de Cursos Pregrabados On-Demand (Competidores directos: Udemy, Domestika).", "Pilar 2: ")
    add_bullet("3. Mentoría Profesional de Alto Valor y Planes Recurrentes (Competidores directos: MentorCruise, ADPList).", "Pilar 3: ")
    add_bullet("4. Red Social Temática, Eventos Grupales y Comunidad (Competidor directo: Skool).", "Pilar 4: ")

    add_callout(
        doc,
        "El objetivo de este Estudio de Mercado no es únicamente descriptivo, sino estrictamente operativo: identificar las "
        "mejores herramientas, micro-interacciones, flujos de pago, sistemas de seguridad y funciones de IA de cada uno de estos "
        "7 gigantes tecnológicos para implementarlos de forma nativa en la reestructuración de CoachingFromHome.",
        title="OBJETIVO ESTRATÉGICO DEL ESTUDIO"
    )

    # =========================================================================
    # SECCIÓN 2: BENCHMARKING DE MARKETPLACES 1 A 1 (PREPLY & SUPERPROF)
    # =========================================================================
    add_h1("2. Benchmarking de Clases y Mentorías 1 a 1: Preply y Superprof")
    add_p(
        "Preply (preply.com) y Superprof (superprof.mx) dominan el mercado mundial de reservas de sesiones 1 a 1 en vivo. "
        "Su éxito radica en eliminar la fricción entre el descubrimiento del mentor y el pago de la primera sesión, así como en "
        "impedir que el profesor y el alumno abandonen la plataforma para evitar pagar comisiones."
    )

    add_h2("2.1 Anatomía de la Tarjeta de Coach de Alta Conversión (Preply Tutor Card)")
    add_p("A diferencia de las tarjetas estáticas tradicionales, Preply utiliza un diseño horizontal de 3 bloques con un panel interactivo:")
    add_bullet("Bloque Izquierdo (Identidad Visual y Presencia): Fotografía cuadrada de bordes redondeados, punto verde pulsante si el coach está conectado en ese instante, bandera del país y etiquetas de idiomas.", "• ")
    add_bullet("Bloque Central (Autoridad y Prueba Social): Incluye insignias verificadas ('Super Coach' para aquellos con calificación >= 4.8 y alta recurrencia; 'Certificado Verificado' cuando el equipo aprobó sus diplomas). Muestra un titular fuerte orientado a resultados, 2 líneas de biografía y un indicador de urgencia en vivo (ej. '6 sesiones reservadas en las últimas 48 horas').", "• ")
    add_bullet("Bloque Derecho (Precio, Rating y CTAs): Precio claro por hora o sesión, calificación promedio con número de reseñas y dos botones: botón primario sólido 'Reservar sesión de prueba' y botón secundario 'Enviar mensaje'.", "• ")
    add_bullet("Panel Flotante 'Split Video Hover' (La herramienta clave a copiar): Cuando el usuario pasa el mouse sobre la tarjeta de cualquier coach, aparece inmediatamente a la derecha un recuadro flotante (sticky) que reproduce el video de presentación del coach (de 60 a 90 segundos) junto con una mini-tabla visual de su disponibilidad semanal dividida en Mañana, Tarde y Noche. El alumno puede ver al coach hablar y comprobar si sus horarios coinciden sin salir del catálogo.", "• ")

    add_h2("2.2 Estrategia de Adquisición de Superprof: '1ª Sesión Gratis' y Velocidad de Respuesta")
    add_bullet("Insignia '1ª Sesión Gratis (15 min)': En Superprof, más del 90% de los mentores ofrecen una primera toma de contacto gratuita de 15 a 30 minutos para evaluar objetivos y generar confianza. En CFH ya contamos con el campo 'freeIntroEnabled' (15 min); debemos destacarlo con un badge visual brillante en la tarjeta del catálogo.", "1. ")
    add_bullet("Métrica de Tiempo de Respuesta: Superprof mide cuánto tarda cada coach en contestar su primer mensaje y muestra públicamente: 'Responde en menos de 1 hora'. Esto presiona positivamente al coach a atender rápido a los prospectos.", "2. ")

    add_h2("2.3 Herramientas del Aula Virtual ('Preply Space') que CFH debe Adoptar")
    add_p("Preply reemplazó los enlaces externos de Zoom/Google Meet por su propia aula integrada en el navegador para controlar toda la experiencia:")
    aula_table = doc.add_table(rows=1, cols=3)
    format_table(
        aula_table,
        col_widths=[1.8, 2.2, 2.5],
        headers=["Herramienta del Aula", "Cómo funciona en Preply", "Cómo implementarla en CFH"],
        rows_data=[
            [
                "Videollamada WebRTC con Token",
                "Llamada integrada en el navegador. Registra el segundo exacto en que entran coach y alumno.",
                "Usar nuestra ruta /api/video/token con JWT efímero (límite estricto de 2 personas en 1-a-1) y auto-confirmar asistencia para liberar el pago."
            ],
            [
                "Pizarra Colaborativa (Whiteboard)",
                "Lienzo compartido en tiempo real donde ambos pueden dibujar, subir PDFs, imágenes y usar puntero láser.",
                "Integrar componente de pizarra tipo Excalidraw/Tldraw sincronizado por WebSockets (Pusher) dentro de /[locale]/sesion/[id]."
            ],
            [
                "Panel Lateral de Notas y Acuerdos",
                "Editor de texto al lado del video donde el tutor apunta vocabulario, enlaces, correcciones y tareas.",
                "Añadir pestaña 'Notas de la Sesión' junto al video que guarde automáticamente los apuntes en la base de datos al terminar la llamada."
            ],
            [
                "Resumen Automático con IA (AI Insights)",
                "Al colgar, la IA transcribe la llamada y genera: resumen de temas, % de tiempo hablado por cada uno y ejercicios sugeridos.",
                "Conectar Whisper + Microservicio LLM (Hermes/OpenAI) al finalizar la sesión para crear un registro en la tabla 'AiArtifact' descargable en PDF."
            ]
        ]
    )

    add_h2("2.4 Sistemas Anti-Desintermediación (Cómo evitar que se pasen el contacto)")
    add_p(
        "Tanto Preply como Superprof combaten la fuga de usuarios combinando barreras técnicas con incentivos económicos:"
    )
    add_bullet("Barrera Técnica en el Chat (Pre-Pago): Un filtro de expresiones regulares y procesamiento de lenguaje natural (NLP) analiza cada mensaje enviado antes de que exista una reserva confirmada. Si detecta números de teléfono (en dígitos '551234...' o en texto 'cinco cinco uno...'), correos ('@', 'arroba', 'gmail'), o palabras como 'WhatsApp', 'Telegram', 'Zoom', 'transferencia directa', bloquea el envío del mensaje, muestra una alerta de seguridad y registra un warning en el historial del usuario.", "A. ")
    add_bullet("Incentivo de Comisión Progresiva y Posicionamiento: En Preply, el coach que mantiene a sus alumnos pagando recurrentemente dentro de la plataforma ve reducida la comisión de la plataforma y recibe el sello 'Super Coach', lo que lo coloca en las primeras posiciones del catálogo. Si lleva al alumno por fuera, pierde horas registradas y cae al fondo de los resultados de búsqueda.", "B. ")
    add_bullet("Garantía de Custodia (Escrow): Se informa claramente al alumno que solo los pagos hechos dentro de la plataforma cuentan con 'Garantía de Reemplazo o Reembolso' en caso de que el coach no se presente.", "C. ")

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 3: BENCHMARKING DE CATÁLOGOS DE CURSOS (UDEMY & DOMESTIKA)
    # =========================================================================
    add_h1("3. Benchmarking de Cursos Pregrabados (LMS): Udemy y Domestika")
    add_p(
        "Para cumplir con la recomendación del experto en ciberseguridad de hacer las tarjetas de cursos más compactas "
        "y desplegar una pestaña emergente al pasar el mouse (Hover Preview), analizamos los dos mejores catálogos del mundo: "
        "Udemy (udemy.com) y Domestika (domestika.org)."
    )

    add_h2("3.1 Diseño Compacto de Grilla y 'Hover Preview Card' (Estilo Udemy)")
    add_p("En Udemy, una pantalla de laptop estándar muestra una cuadrícula limpia de 4 a 5 cursos por fila con imágenes compactas (relación 16:9 de 240x135 px). Cuando el usuario posa el cursor sobre cualquier tarjeta durante 300 milisegundos, se despliega un Popover interactivo con la siguiente anatomía exacta:")
    add_bullet("1. Título completo del curso y badge de categoría o 'Bestseller'.", "• ")
    add_bullet("2. Fecha de última actualización (ej. 'Actualizado en Octubre 2026') + Total de horas de video + Nivel de dificultad + Idioma/Subtítulos.", "• ")
    add_bullet("3. Resumen ejecutivo en 2 líneas de lo que trata el curso.", "• ")
    add_bullet("4. Lista de 3 a 4 puntos clave ('Lo que aprenderás') precedidos por un icono de checkmark (✓), extraídos del campo keyPoints del curso.", "• ")
    add_bullet("5. Botón principal de acción inmediata: 'Añadir al carrito' en 1 solo clic + Botón de guardar en favoritos (corazón).", "• ")

    add_h2("3.2 Autoplay Silencioso en Miniatura y Galería de Proyectos (Estilo Domestika)")
    add_bullet("Video Tráiler en Hover (Muted Autoplay): En Domestika, al pasar el cursor sobre la tarjeta del curso, la miniatura estática se reemplaza suavemente por el video introductorio del curso reproduciéndose sin sonido (con un botón para activar el audio). Esto permite evaluar la calidad de producción del curso en 3 segundos sin cambiar de página.", "1. ")
    add_bullet("Porcentaje de Valoraciones Positivas: En lugar de solo estrellas, Domestika muestra un indicador de confianza directa: '99% de valoraciones positivas (1,240 alumnos)'.", "2. ")
    add_bullet("Galería Pública de Proyectos Finales: Cada curso tiene una pestaña donde los alumnos que terminaron el curso publican las fotos o documentos de su proyecto final. Cuando un comprador potencial ve los trabajos reales creados por otros alumnos, la tasa de conversión de compra se multiplica.", "3. ")

    add_h2("3.3 Experiencia de Aula de Curso Pregrabado (Reproductor y Notas por Minuto)")
    add_bullet("Vista Previa Gratuita de Lecciones Seleccionadas: Permitir que el coach marque 1 o 2 lecciones de su curso con el flag 'isPreview = true' para que cualquier visitante pueda verlas gratis desde el temario antes de comprar.", "• ")
    add_bullet("Notas Ancladas al Segundo del Video ([mm:ss]): Nuestra base de datos ya cuenta con el modelo 'LessonNote' con el campo 'seconds'. Al igual que en Udemy, debemos mostrar debajo del video un botón 'Crear nota en 04:15'. Al hacer clic en esa nota días después, el reproductor de Bunny.net salta automáticamente al segundo 255.", "• ")
    add_bullet("Protección Total Anti-Descarga: Al igual que Udemy Business y Domestika, los videos de cursos se sirven por streaming HLS segmentado (.m3u8) con tokens firmados, deshabilitando el clic derecho, la impresión y la descarga directa del archivo original.", "• ")

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 4: BENCHMARKING DE AUTORIDAD Y MENTORÍA EJECUTIVA (MENTORCRUISE & ADPLIST)
    # =========================================================================
    add_h1("4. Benchmarking de Autoridad, Credibilidad y Planes: MentorCruise y ADPList")
    add_p(
        "En los audios de la sesión de trabajo se enfatizó un punto crucial: 'Tenemos que vender credibilidad real; que la persona "
        "que entre vea que el coach es una persona real con trayectoria comprobable, apariciones en medios, entrevistas y resultados'. "
        "MentorCruise (mentorcruise.com) y ADPList (adplist.org) son los referentes mundiales en construir autoridad para mentores."
    )

    add_h2("4.1 Cómo Presentar la Credibilidad y Autoridad del Coach (Social Proof)")
    add_bullet("Contador Auditable de Minutos de Mentoría (Estilo ADPList): El perfil del coach calcula y exhibe automáticamente el total de minutos de sesiones en vivo que ha impartido dentro de la plataforma (ej. '3,840 minutos de coaching impartidos · 64 sesiones completadas · 99% de puntualidad'). Esta métrica no se puede inventar y genera confianza inmediata.", "1. ")
    add_bullet("Sección 'Apariciones en Medios, Prensa y Credenciales' (Due Diligence): Añadir en el perfil del coach un bloque estructurado donde pueda enlazar: entrevistas en televisión/radio, artículos publicados, canal oficial de YouTube, perfil de LinkedIn verificado y certificaciones validadas por nuestro motor OCR.", "2. ")
    add_bullet("Insignias de Trayectoria ('Top 1% Coach' / 'Coach Fundador'): Reconocimientos visuales automáticos otorgados por el sistema según las reseñas de 5 estrellas y el número de alumnos graduados.", "3. ")
    add_bullet("Fotografía Real Verificada: Prohibición de avatares generados por IA o ilustraciones genéricas en perfiles de coach; validación de fotografía real frente a documento o video de bienvenida obligatorio.", "4. ")

    add_h2("4.2 Modelo de Monetización Híbrido: Planes Mensuales vs. Sesiones Sueltas")
    add_p(
        "MentorCruise demostró que vender únicamente horas sueltas limita las ganancias del coach y de la plataforma. "
        "Su solución es ofrecer dos pestañas de compra dentro del perfil de cada coach:"
    )
    add_bullet("Opción A - Sesiones Individuales (One-Off Sessions): Sesión introductoria (15 min gratis), Sesión de Diagnóstico de 60 min ($X USD) o Paquete de 4 sesiones con 10% de descuento.", "• ")
    add_bullet("Opción B - Plan de Mentoría Mensual (Retainer Suscripción): El alumno paga una mensualidad fija (ej. $120 USD/mes) que le incluye: 2 o 4 videollamadas de 60 min al mes + acceso a consultas por chat asíncrono con respuesta garantizada en menos de 24 horas + seguimiento de tareas.", "• ")

    add_h2("4.3 Sesiones Grupales y Masterclasses en Vivo (ADPList Group Sessions)")
    add_p(
        "Tal como se propuso en el cuarto audio de la reunión (para clases grupales de programación, arte, yoga o negocios), "
        "ADPList permite a los mentores abrir 'Group Sessions' con las siguientes características:"
    )
    add_bullet("El coach define título, fecha, hora, precio por boleto (o gratuito promocional) y cupo máximo estricto (ej. 10, 20 o 50 asientos).", "1. ")
    add_bullet("La tarjeta del evento muestra los avatares de los alumnos ya inscritos y un indicador de escasez: '¡Últimos 3 lugares disponibles de 15!'.", "2. ")
    add_bullet("Al terminar la sesión grupal, el sistema envía automáticamente a todos los asistentes un enlace para reservar una sesión privada 1 a 1 con ese coach, funcionando como una máquina de ventas.", "3. ")

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 5: BENCHMARKING DE COMUNIDAD Y GAMIFICACIÓN (SKOOL)
    # =========================================================================
    add_h1("5. Benchmarking de Comunidad, Eventos y Gamificación: Skool")
    add_p(
        "Skool (skool.com) revolucionó las plataformas educativas al transformar los foros aburridos en comunidades gamificadas "
        "donde los usuarios entran todos los días. A continuación se detallan las mecánicas de Skool que podemos incorporar en la "
        "sección 'Comunidad' de CoachingFromHome:"
    )

    add_h2("5.1 Calendario de Eventos en Vivo Integrado en la Comunidad")
    add_bullet("En la parte superior de la Comunidad se añade la pestaña 'Calendario de Eventos'.", "• ")
    add_bullet("Muestra en una cuadrícula mensual todas las sesiones grupales, masterclasses y talleres en vivo organizados por los coaches de la plataforma, convirtiendo automáticamente la hora a la zona horaria local del estudiante.", "• ")
    add_bullet("Permite inscribirse en 1 clic y agregar el recordatorio a Google Calendar, Apple Calendar o Outlook.", "• ")

    add_h2("5.2 Gamificación por Aportes de Calidad (1 Like = 1 Punto)")
    add_bullet("Mecánica Anti-Spam: El usuario NO gana puntos por publicar muchos mensajes vacíos; solo gana 1 punto cuando otro miembro de la comunidad le da 'Like' a su publicación o comentario porque le resultó útil.", "1. ")
    add_bullet("Sistema de 9 Niveles de Prestigio: El avatar de cada alumno y coach muestra un badge con su nivel actual (Nivel 1 a Nivel 9) basado en sus puntos acumulados.", "2. ")
    add_bullet("Tabla de Clasificación (Leaderboard): Panel lateral en la comunidad que muestra a los miembros y coaches que más valor han aportado en los últimos 7 días, 30 días e histórico.", "3. ")
    add_bullet("Desbloqueo de Recompensas por Nivel: Al alcanzar por ejemplo el Nivel 3 (20 puntos), el estudiante desbloquea automáticamente un cupón de descuento para su próxima sesión 1 a 1 o acceso a recursos exclusivos.", "4. ")

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 6: MATRIZ COMPARATIVA MAESTRA Y PLAN DE ADOPCIÓN PARA CFH
    # =========================================================================
    add_h1("6. Matriz Comparativa Maestra y Plan de Adopción para CFH")
    add_p(
        "La siguiente tabla resume las funcionalidades exactas que CoachingFromHome tomará de cada referente del mercado "
        "para consolidar una plataforma superior:"
    )

    master_table = doc.add_table(rows=1, cols=4)
    format_table(
        master_table,
        col_widths=[1.2, 1.5, 1.9, 1.9],
        headers=["Plataforma", "Área de Excelencia", "Funcionalidad / Herramienta a Copiar", "Implementación Técnica en CFH"],
        rows_data=[
            [
                "Preply",
                "Catálogo de Coaches y Aula 1-a-1",
                "Split Video Hover + Mini horario + Resumen IA post-llamada + Pizarra.",
                "Panel lateral en hover en /[locale]/coaches y pipeline Whisper + LLM al cerrar sala de video."
            ],
            [
                "Superprof",
                "Conversión de Primera Reserva",
                "Badge '1ª Sesión Gratis (15 min)' + Tiempo promedio de respuesta + Filtro en chat.",
                "Destacar freeIntroEnabled en UI, calcular tiempo de respuesta en Conversation y filtro Regex en chat."
            ],
            [
                "Udemy",
                "Catálogo de Cursos (LMS)",
                "Tarjetas compactas (4 por fila) + Popover Hover Card con 'Lo que aprenderás' y Carrito 1-clic.",
                "Rediseñar CourseCard con popover flotante al pasar el mouse y botón directo a CartContext."
            ],
            [
                "Domestika",
                "Atractivo Visual y Prueba Social",
                "Autoplay silencioso de tráiler en hover + Galería de proyectos subidos por graduados.",
                "Reproducir introVideoUrl de Bunny.net en muted hover y mostrar entregas destacadas."
            ],
            [
                "MentorCruise",
                "Ingresos Recurrentes",
                "Planes mensuales de acompañamiento (Suscripción al Coach: llamadas + chat).",
                "Añadir modalidad de paquete mensual recurrente en el perfil del coach vía Stripe Subscriptions."
            ],
            [
                "ADPList",
                "Autoridad y Eventos Grupales",
                "Contador de minutos totales impartidos + Sección Prensa/Medios + Sesiones Grupales.",
                "Sumar minutos de Bookings completados en perfil, sección de prensa y llamadas con aforo N."
            ],
            [
                "Skool",
                "Comunidad y Retención",
                "Calendario mensual de eventos en vivo + Gamificación (1 Like = 1 Punto) + Leaderboard.",
                "Añadir vista de calendario en /[locale]/comunidad y conteo de CommunityLikes para niveles."
            ]
        ]
    )

    add_h2("6.1 Conclusión Estratégica para el Equipo de Desarrollo y Ciberseguridad")
    add_p(
        "Al combinar la solidez de ciberseguridad (WAF, Cloudflare Family DNS 1.1.1.3, autenticación OTP de 6 dígitos, "
        "tokens JWT efímeros en videollamadas, almacenamiento efímero de grabaciones y microservicios de IA open-source) "
        "con los patrones de conversión probados de Preply, Udemy, Domestika, MentorCruise y Skool, CoachingFromHome "
        "se posicionará no solo como una página web funcional, sino como una plataforma EdTech de élite altamente rentable, "
        "segura y escalable."
    )

    output_path_1 = os.path.join(os.path.dirname(__file__), "Estudio de Mercado.docx")
    output_path_2 = os.path.join(os.path.dirname(__file__), "Estudio_de_Mercado_CFH.docx")
    doc.save(output_path_1)
    doc.save(output_path_2)
    print(f"Market study documents created successfully at:\n1) {output_path_1}\n2) {output_path_2}")

if __name__ == "__main__":
    create_market_study_doc()
